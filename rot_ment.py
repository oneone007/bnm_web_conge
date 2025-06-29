import oracledb
from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
import logging
import pandas as pd
import io
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import datetime

app = Flask(__name__)
CORS(app)  # Enable CORS for all routes

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Configure Oracle database connection pool
DB_POOL = oracledb.create_pool(
    user="compiere",
    password="compiere",
    dsn="192.168.1.213/compiere",
    min=2,
    max=10,
    increment=1
)


def fetch_fournisseur_by_product(product):
    try:
        # Acquire a connection from the pool
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()

            query = """
                SELECT DISTINCT
                    CAST(cb.name AS VARCHAR2(300)) AS FOURNISSEUR
                FROM 
                    M_InOut xf
                JOIN M_INOUTLINE mi ON mi.M_INOUT_ID = xf.M_INOUT_ID
                JOIN C_BPartner cb ON cb.C_BPARTNER_ID = xf.C_BPARTNER_ID
                JOIN M_PRODUCT m ON m.M_PRODUCT_id = mi.M_PRODUCT_id
                WHERE 
                    UPPER(m.name) LIKE UPPER(:product) || '%'
                    AND xf.AD_Org_ID = 1000000
                    AND xf.C_DocType_ID IN (1000013, 1000646)
                    AND xf.M_Warehouse_ID IN (1000724, 1000000, 1000720, 1000725)
                ORDER BY 
                    FOURNISSEUR
            """

            params = {
                'product': product or None
            }

            # Execute the query with the provided parameters
            cursor.execute(query, params)

            # Fetch the results
            rows = cursor.fetchall()

            # Format the results into a simple list of supplier names
            suppliers = [row[0] for row in rows]

            return suppliers
    
    except Exception as e:
        logger.error(f"Error fetching suppliers by product: {e}")
        return {"error": "An error occurred while fetching suppliers by product."}

# Flask route to handle the request
@app.route('/fetchSuppliersByProduct', methods=['GET'])
def fetch_suppliers_by_product():
    product = request.args.get('product')

    # Ensure product is provided
    if not product:
        return jsonify({"error": "Missing product parameter"}), 400

    # Fetch data from the database
    suppliers = fetch_fournisseur_by_product(product)

    # Return the result as a JSON response
    return jsonify(suppliers)




@app.route('/rotation_monthly_achat', methods=['GET'])
def fetch_rotation_achat():
    years = request.args.get('years')  # Comma-separated list of years
    fournisseur_param = request.args.get('fournisseur')  # Can be comma-separated
    product = request.args.get('product')

    # If years is not provided, use current year
    if not years:
        current_year = datetime.now().year
        years = str(current_year)
    
    # Handle multiple suppliers (comma-separated or single)
    if fournisseur_param:
        fournisseurs = [f.strip() for f in fournisseur_param.split(',') if f.strip()]
    else:
        return jsonify({"error": "At least one supplier is required"}), 400

    try:
        # Split the years string into a list and convert to integers
        year_list = [int(year.strip()) for year in years.split(',')]
    except ValueError:
        return jsonify({"error": "Invalid years format. Please provide comma-separated years (e.g., 2022,2023,2024)"}), 400

    # Fetch data from the database with the list of suppliers
    data = fetch_rotation_monthly_achat(year_list, fournisseurs, product)

    # Return the result as a JSON response
    return jsonify(data)


def fetch_rotation_monthly_achat(year_list, fournisseurs, product=None):
    try:
        # Acquire a connection from the pool
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()

            # Build the supplier condition
            if isinstance(fournisseurs, str):
                fournisseurs = [fournisseurs]

            supplier_conditions = []
            for i, f in enumerate(fournisseurs):
                bind_var = f'fournisseur_{i}'
                supplier_conditions.append(f"UPPER(cb.name) LIKE UPPER(:{bind_var}) || '%'")

            supplier_condition = ' OR '.join(supplier_conditions) if supplier_conditions else '1=0'

            # Create the years parameter list
            year_placeholders = ', '.join([f':year_{i}' for i in range(len(year_list))])
            
            query = """
                WITH monthly_data AS (
                    SELECT 
                        TO_CHAR(xf.MOVEMENTDATE, 'YYYY') AS year,
                        TO_CHAR(xf.MOVEMENTDATE, 'MM') AS month_num,
                        TO_CHAR(xf.MOVEMENTDATE, 'Month') AS month_name,
                        m.name AS produit,
                        cb.name AS fournisseur,
                        SUM(CASE 
                            WHEN xf.C_DocType_ID = 1000646 THEN -1 * TO_NUMBER(mi.QTYENTERED)
                            ELSE TO_NUMBER(mi.QTYENTERED)
                        END) AS qty,
                        SUM(CASE 
                            WHEN xf.C_DocType_ID = 1000646 THEN -1 * TO_NUMBER(ma.valuenumber) * TO_NUMBER(mi.QTYENTERED)
                            ELSE TO_NUMBER(ma.valuenumber) * TO_NUMBER(mi.QTYENTERED)
                        END) AS chiffre
                    FROM 
                        M_InOut xf
                        JOIN M_INOUTLINE mi ON mi.M_INOUT_ID = xf.M_INOUT_ID
                        JOIN M_ATTRIBUTEINSTANCE ma ON ma.M_ATTRIBUTESETINSTANCE_ID = mi.M_ATTRIBUTESETINSTANCE_ID
                        JOIN C_BPartner cb ON cb.C_BPARTNER_ID = xf.C_BPARTNER_ID
                        JOIN M_PRODUCT m ON m.M_PRODUCT_id = mi.M_PRODUCT_id
                    WHERE 
                        TO_CHAR(xf.MOVEMENTDATE, 'YYYY') IN (""" + year_placeholders + """)
                        AND xf.AD_Org_ID = 1000000
                        AND xf.C_DocType_ID IN (1000013, 1000646)
                        AND ma.M_Attribute_ID = 1000504
                        AND xf.DOCSTATUS = 'CO'
                        AND xf.M_Warehouse_ID IN (1000724, 1000000, 1000720, 1000725)
                        AND (""" + supplier_condition + """)
                        AND (:product IS NULL OR UPPER(m.name) LIKE UPPER(:product) || '%')
                    GROUP BY 
                        TO_CHAR(xf.MOVEMENTDATE, 'YYYY'),
                        TO_CHAR(xf.MOVEMENTDATE, 'MM'),
                        TO_CHAR(xf.MOVEMENTDATE, 'Month'),
                        m.name,
                        cb.name
                )
                SELECT * FROM monthly_data
                ORDER BY year, month_num, produit
            """

            # Prepare parameters
            params = {'product': product}
            
            # Add supplier parameters
            for i, f in enumerate(fournisseurs):
                params[f'fournisseur_{i}'] = f

            # Add year parameters
            for i, year in enumerate(year_list):
                params[f'year_{i}'] = str(year)

            # Execute the query with the parameters
            cursor.execute(query, params)

            # Fetch the results
            rows = cursor.fetchall()

            # Format the results into a nested dictionary structure
            result = {}
            
            # Process each row and organize by year and month
            for row in rows:
                year, month_num, month_name, produit, fournisseur, qty, chiffre = row
                month_num = month_num.zfill(2)  # Ensure two digits for month number
                
                if year not in result:
                    result[year] = {}
                
                if month_num not in result[year]:
                    result[year][month_num] = {
                        'month_name': month_name.strip(),
                        'details': [],
                        'total': {'QTY': 0, 'CHIFFRE': 0}
                    }
                
                # Add product details with supplier info
                result[year][month_num]['details'].append({
                    'PRODUIT': produit,
                    'FOURNISSEUR': fournisseur,
                    'QTY': float(qty) if qty is not None else 0,
                    'CHIFFRE': float(chiffre) if chiffre is not None else 0
                })
                
                # Update monthly totals
                result[year][month_num]['total']['QTY'] += float(qty) if qty is not None else 0
                result[year][month_num]['total']['CHIFFRE'] += float(chiffre) if chiffre is not None else 0

            # Sort years and months numerically
            sorted_result = {
                str(year): dict(sorted(months.items()))
                for year, months in sorted(result.items())
            }
            return sorted_result

    except Exception as e:
        logger.error(f"Error fetching product recap achat: {e}")
        return {"error": "An error occurred while fetching product recap achat."}


from flask import make_response
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch
from datetime import datetime
import io

# Define a color palette for years (can be extended as needed)
YEAR_COLORS = [
    colors.HexColor('#4E79A7'),  # Blue
    colors.HexColor('#F28E2B'),  # Orange
    colors.HexColor('#E15759'),  # Red
    colors.HexColor('#76B7B2'),  # Teal
    colors.HexColor('#59A14F'),  # Green
    colors.HexColor('#EDC948'),  # Yellow
    colors.HexColor('#B07AA1'),  # Purple
    colors.HexColor('#FF9DA7'),  # Pink
    colors.HexColor('#9C755F'),  # Brown
    colors.HexColor('#BAB0AC')   # Gray
]

@app.route('/rotation_monthly_achat_pdf', methods=['GET'])
def download_product_achat_pdf():
    from flask import request, jsonify

    years = request.args.get('years')
    fournisseur_param = request.args.get('fournisseur')
    product = request.args.get('product')

    if not years:
        years = str(datetime.now().year)

    try:
        year_list = [int(y.strip()) for y in years.split(',')]
    except ValueError:
        return jsonify({"error": "Invalid years format"}), 400

    # Handle multiple suppliers (comma-separated or single)
    if fournisseur_param:
        fournisseurs = [f.strip() for f in fournisseur_param.split(',') if f.strip()]
    else:
        return jsonify({"error": "Fournisseur parameter is required"}), 400

    data = fetch_rotation_monthly_achat(year_list, fournisseurs, product)
    if 'error' in data:
        return jsonify(data), 500

    # Create descriptive filename
    supplier_text = f"{len(fournisseurs)}_suppliers" if len(fournisseurs) > 1 else fournisseurs[0][:20]
    filename = f"achat_recap_{years.replace(',', '-')}_{supplier_text}_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=20, leftMargin=20, topMargin=30, bottomMargin=30)
    elements = []

    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    normal_style = styles["Normal"]
    table_total_style = ParagraphStyle("row", parent=normal_style, alignment=1, textColor=colors.blue, fontName="Helvetica-Bold", fontSize=6)

    table_header_style = ParagraphStyle("header", parent=normal_style, alignment=1, textColor=colors.white, fontSize=8)
    table_row_style = ParagraphStyle("row", parent=normal_style, alignment=1, fontSize=6)

    elements.append(Paragraph("Récapitulatif des Achats", title_style))
    
    # Create supplier display text
    supplier_display = ", ".join(fournisseurs) if len(fournisseurs) <= 3 else f"{len(fournisseurs)} suppliers"
    elements.append(Paragraph(f"Fournisseurs: {supplier_display} | Produit: {product or 'Tous'} | Années: {years}", styles["Heading2"]))
    elements.append(Spacer(1, 0.25 * inch))

    for year_idx, (year, months_data) in enumerate(data.items()):
        year_color = YEAR_COLORS[year_idx % len(YEAR_COLORS)]
        year_style = ParagraphStyle("year", parent=styles["Heading1"], textColor=colors.white, backColor=year_color)
        elements.append(Paragraph(f"Année: {year}", year_style))

        for part_label, month_range in [("Janvier à Juin", range(1, 7)), ("Juillet à Décembre", range(7, 13))]:
            # For multiple suppliers, create a combined product map with supplier info
            if len(fournisseurs) > 1:
                combined_products = {}
                for m in month_range:
                    m_str = str(m).zfill(2)
                    if m_str in months_data:
                        for item in months_data[m_str]["details"]:
                            product_key = f"{item['PRODUIT']} ({item['FOURNISSEUR']})"
                            if product_key not in combined_products:
                                combined_products[product_key] = {}
                            combined_products[product_key][m] = item
                unique_products = list(combined_products.keys())
            else:
                # Single supplier - use original logic
                unique_products = set()
                for m in month_range:
                    m_str = str(m).zfill(2)
                    if m_str in months_data:
                        for item in months_data[m_str]["details"]:
                            unique_products.add(item["PRODUIT"])
                unique_products = list(unique_products)

            if not unique_products:
                continue

            # Header rows
            header_row1 = [Paragraph("Produit", table_header_style)]
            for m in month_range:
                m_str = str(m).zfill(2)
                month_name = months_data.get(m_str, {}).get("month_name", datetime(2025, m, 1).strftime("%b")).capitalize()
                header_row1.extend([Paragraph(month_name, table_header_style), ""])

            header_row2 = [""]
            for _ in month_range:
                header_row2.extend([
                    Paragraph("Qty", table_header_style),
                    Paragraph("Prix", table_header_style)
                ])

            table_data = [header_row1, header_row2]

            # Product rows
            for prod in sorted(unique_products):
                row = [Paragraph(prod, table_row_style)]
                for m in month_range:
                    m_str = str(m).zfill(2)
                    month_data = months_data.get(m_str, {})
                    
                    if len(fournisseurs) > 1:
                        # For combined view, find the specific product entry
                        detail = combined_products.get(prod, {}).get(m)
                        qty = f"{detail['QTY']:,.0f}" if detail else "-"
                        amt = f"{detail['CHIFFRE']:,.2f}" if detail else "-"
                    else:
                        # Single supplier - use original logic
                        detail = next((d for d in month_data.get("details", []) if d["PRODUIT"] == prod), None)
                        qty = f"{detail['QTY']:,.0f}" if detail else "-"
                        amt = f"{detail['CHIFFRE']:,.2f}" if detail else "-"
                    
                    row.append(Paragraph(qty, table_row_style))
                    row.append(Paragraph(amt, table_row_style))
                table_data.append(row)

            # Total row
            # Total row
            total_row = [Paragraph("TOTAL", table_total_style)]
            for m in month_range:
                m_str = str(m).zfill(2)
                month_data = months_data.get(m_str, {})
                total = month_data.get("total", {})

                qty_val = total.get("QTY")
                amt_val = total.get("CHIFFRE")

                qty = f"{qty_val:,.0f}" if qty_val is not None else "-"
                amt = f"{amt_val:,.2f}" if amt_val is not None else "-"

                total_row.append(Paragraph(qty, table_total_style))
                total_row.append(Paragraph(amt, table_total_style))
            table_data.append(total_row)

            # Table construction
            col_widths = [1.9 * inch] + [0.5 * inch] * (len(month_range) * 2)
            table = Table(table_data, colWidths=col_widths)

            # Create span styles for merged headers
            span_style = []
            for i in range(len(month_range)):
                col_start = 1 + i * 2
                span_style.append(('SPAN', (col_start, 0), (col_start + 1, 0)))

            # Final style
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 1), year_color),
                ('TEXTCOLOR', (0, 0), (-1, 1), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 2),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('ROWBACKGROUNDS', (0, 2), (-1, -2), [colors.whitesmoke, colors.white]),
            ] + span_style))

            elements.append(Paragraph(f"{part_label}", styles["Heading3"]))
            elements.append(table)
            elements.append(Spacer(1, 0.3 * inch))

    doc.build(elements)
    buffer.seek(0)
    response = make_response(buffer.getvalue())
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = f"attachment; filename={filename}"

    return response


@app.route('/rotation_monthly_vente_pdf', methods=['GET'])
def download_product_vente_pdf():
    from flask import request, jsonify

    years = request.args.get('years')
    fournisseur_param = request.args.get('fournisseur')
    product = request.args.get('product')
    client_param = request.args.get('client')
    zone_param = request.args.get('zone')

    if not years:
        years = str(datetime.now().year)

    try:
        year_list = [int(y.strip()) for y in years.split(',')]
    except ValueError:
        return jsonify({"error": "Invalid years format"}), 400

    # Handle multiple suppliers (comma-separated or single)
    fournisseurs = None
    if fournisseur_param:
        fournisseurs = [f.strip() for f in fournisseur_param.split(',') if f.strip()]

    # Handle multiple clients (comma-separated or single)
    clients = None
    if client_param:
        clients = [c.strip() for c in client_param.split(',') if c.strip()]

    # Handle multiple zones (comma-separated or single)
    zones = None
    if zone_param:
        zones = [z.strip() for z in zone_param.split(',') if z.strip()]

    # Handle multiple products (comma-separated or single)
    products = None
    if product:
        products = [p.strip() for p in product.split(',') if p.strip()]

    data = rot_mont_vente(year_list, fournisseurs, products, clients, zones)
    if 'error' in data:
        return jsonify(data), 500

    # Create descriptive filename
    product_text = product or 'all_products'
    filename = f"vente_recap_{years.replace(',', '-')}_{product_text[:20]}_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=20, leftMargin=20, topMargin=30, bottomMargin=30)
    elements = []

    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    normal_style = styles["Normal"]
    table_total_style = ParagraphStyle("row", parent=normal_style, alignment=1, textColor=colors.blue, fontName="Helvetica-Bold", fontSize=6)

    table_header_style = ParagraphStyle("header", parent=normal_style, alignment=1, textColor=colors.white, fontSize=8)
    table_row_style = ParagraphStyle("row", parent=normal_style, alignment=1, fontSize=6)

    elements.append(Paragraph("Récapitulatif des Ventes", title_style))
    
    # Create filter display text
    filter_parts = []
    if clients:
        client_display = ", ".join(clients) if len(clients) <= 3 else f"{len(clients)} clients"
        filter_parts.append(f"Clients: {client_display}")
    else:
        filter_parts.append("Clients: Tous")
    
    if fournisseurs:
        supplier_display = ", ".join(fournisseurs) if len(fournisseurs) <= 3 else f"{len(fournisseurs)} fournisseurs"
        filter_parts.append(f"Fournisseurs: {supplier_display}")
    
    if zones:
        zone_display = ", ".join(zones) if len(zones) <= 3 else f"{len(zones)} zones"
        filter_parts.append(f"Zones: {zone_display}")
    
    filter_parts.append(f"Produit: {product or 'Tous'}")
    filter_parts.append(f"Années: {years}")
    
    elements.append(Paragraph(" | ".join(filter_parts), styles["Heading2"]))
    elements.append(Spacer(1, 0.25 * inch))

    for year_idx, (year, months_data) in enumerate(data.items()):
        year_color = YEAR_COLORS[year_idx % len(YEAR_COLORS)]
        year_style = ParagraphStyle("year", parent=styles["Heading1"], textColor=colors.white, backColor=year_color)
        elements.append(Paragraph(f"Année: {year}", year_style))

        for part_label, month_range in [("Janvier à Juin", range(1, 7)), ("Juillet à Décembre", range(7, 13))]:
            # Create a combined product map organized by product and supplier (not client)
            product_suppliers = {}
            for m in month_range:
                m_str = str(m).zfill(2)
                if m_str in months_data:
                    for item in months_data[m_str]["details"]:
                        product_name = item['PRODUIT']
                        supplier_name = item.get('FOURNISSEUR', 'Unknown Supplier')
                        
                        # Create product-supplier key
                        product_key = f"{product_name} ({supplier_name})"
                        
                        if product_key not in product_suppliers:
                            product_suppliers[product_key] = {}
                        
                        # Aggregate data for the same product-supplier-month combination
                        if m not in product_suppliers[product_key]:
                            product_suppliers[product_key][m] = {
                                'QTY': 0,
                                'TOTAL': 0,
                                'CONSOMATION': 0
                            }
                        
                        product_suppliers[product_key][m]['QTY'] += item['QTY']
                        product_suppliers[product_key][m]['TOTAL'] += item['TOTAL']
                        # We need consomation to calculate margin properly
                        # For now, we'll calculate margin from the data we have
                        item_marge = item['MARGE'] if item['MARGE'] is not None else 0
                        consomation = item['TOTAL'] / (1 + item_marge) if item_marge > 0 else item['TOTAL']
                        product_suppliers[product_key][m]['CONSOMATION'] += consomation

            if not product_suppliers:
                continue

            # Create table for this period
            elements.append(Paragraph(f"{part_label}", styles["Heading3"]))
            
            # Header row - single row with month names
            header_row = [Paragraph("Produit (Fournisseur)", table_header_style)]
            for m in month_range:
                m_str = str(m).zfill(2)
                month_name = months_data.get(m_str, {}).get("month_name", datetime(2025, m, 1).strftime("%b")).capitalize()
                header_row.append(Paragraph(month_name, table_header_style))

            table_data = [header_row]

            # Product-supplier rows
            period_total_qty = {m: 0 for m in month_range}
            period_total_amt = {m: 0 for m in month_range}
            period_total_consomation = {m: 0 for m in month_range}
            
            for product_key in sorted(product_suppliers.keys()):
                row = [Paragraph(product_key, table_row_style)]
                for m in month_range:
                    detail = product_suppliers.get(product_key, {}).get(m)
                    if detail:
                        qty = f"{detail['QTY']:,.0f}"
                        total = f"{detail['TOTAL']:,.2f}"
                        # Calculate margin percentage
                        if detail['CONSOMATION'] > 0:
                            marge_pct = ((detail['TOTAL'] - detail['CONSOMATION']) / detail['CONSOMATION']) * 100
                            marge = f"{marge_pct:.2f}%"
                        else:
                            marge = "0.00%"
                        
                        # Create vertical layout like the frontend with better formatting
                        month_data_parts = [
                            f"<b>Qty:</b> {qty}",
                            f"<b>Total:</b> {total}",
                            f"<b>Marge:</b> {marge}"
                        ]
                        month_data = "<br/>".join(month_data_parts)
                        
                        period_total_qty[m] += detail['QTY']
                        period_total_amt[m] += detail['TOTAL']
                        period_total_consomation[m] += detail['CONSOMATION']
                    else:
                        month_data = "<b>Qty:</b> -<br/><b>Total:</b> -<br/><b>Marge:</b> -"
                    
                    row.append(Paragraph(month_data, table_row_style))
                table_data.append(row)

            # Period total row
            total_row = [Paragraph("TOTAL", table_total_style)]
            for m in month_range:
                qty = f"{period_total_qty[m]:,.0f}" if period_total_qty[m] > 0 else "0"
                amt = f"{period_total_amt[m]:,.2f}" if period_total_amt[m] > 0 else "0"
                
                if period_total_consomation[m] > 0:
                    total_marge_pct = ((period_total_amt[m] - period_total_consomation[m]) / period_total_consomation[m]) * 100
                    marge = f"{total_marge_pct:.2f}%"
                else:
                    marge = "0.00%"
                
                # Create vertical layout for totals too with better formatting
                total_data_parts = [
                    f"<b>Qty:</b> {qty}",
                    f"<b>Total:</b> {amt}",
                    f"<b>Marge:</b> {marge}"
                ]
                total_data = "<br/>".join(total_data_parts)
                total_row.append(Paragraph(total_data, table_total_style))
            table_data.append(total_row)

            # Table construction - single column per month now
            col_widths = [2.0 * inch] + [1.0 * inch] * len(month_range)
            table = Table(table_data, colWidths=col_widths, repeatRows=1)

            # Simpler style without span since we have single column per month
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), year_color),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),  # Top alignment for better readability
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 0), (-1, 0), 8),  # Header font size
                ('FONTSIZE', (0, 1), (-1, -2), 7),  # Data rows font size
                ('FONTSIZE', (0, -1), (-1, -1), 8),  # Total row font size
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.whitesmoke, colors.white]),
                ('BACKGROUND', (0, -1), (-1, -1), colors.lightgrey),  # Total row background
                ('TEXTCOLOR', (0, -1), (-1, -1), colors.black),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ('TOPPADDING', (0, 1), (-1, -1), 6),    # More padding for data rows
                ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
                ('LEADING', (0, 1), (-1, -1), 12),       # Line spacing
            ]))

            elements.append(table)
            elements.append(Spacer(1, 0.2 * inch))

    doc.build(elements)
    buffer.seek(0)
    response = make_response(buffer.getvalue())
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = f"attachment; filename={filename}"

    return response


@app.route('/fetchFournisseurRecapAchatByYear', methods=['GET'])
def fetch_fournisseur_recap_achat_by_year():
    fournisseur = request.args.get('fournisseur')
    product = request.args.get('product')
    years = request.args.getlist('years')  # Get all years parameters
    month = request.args.get('month')

    try:
        years = [int(y) for y in years] if years else None
        month = int(month) if month else None
    except ValueError:
        return jsonify({"error": "Invalid year or month format"}), 400

    if not years:
        # If no years specified, get current year
        current_year = datetime.now().year
        years = [current_year]

    full_data = {}
    for year in years:
        yearly_data = fetch_fournisseur_recap_achat_data_by_year(
            fournisseur=fournisseur,
            product=product,
            year=year,
            month=month
        )
        full_data[str(year)] = yearly_data

    return jsonify(full_data)




def fetch_fournisseur_recap_achat_data_by_year(fournisseur=None, product=None, year=None, month=None):
    try:
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()

            if year and month:
                start_date = f"{year}-{month:02d}-01"
                last_day = 31 if month in [1, 3, 5, 7, 8, 10, 12] else 30
                if month == 2:
                    last_day = 29 if year % 4 == 0 and (year % 100 != 0 or year % 400 == 0) else 28
                end_date = f"{year}-{month:02d}-{last_day}"
            elif year:
                start_date = f"{year}-01-01"
                end_date = f"{year}-12-31"
            else:
                current_year = datetime.now().year
                start_date = f"{current_year}-01-01"
                end_date = f"{current_year}-12-31"
                year = current_year

            query = """
                SELECT 
                    TO_CHAR(xf.MOVEMENTDATE, 'YYYY') AS year,
                    TO_CHAR(xf.MOVEMENTDATE, 'MM') AS month,
                    SUM(CASE 
                        WHEN xf.C_DocType_ID = 1000646 THEN -1 * TO_NUMBER(ma.valuenumber) * TO_NUMBER(mi.QTYENTERED) 
                        ELSE TO_NUMBER(ma.valuenumber) * TO_NUMBER(mi.QTYENTERED) 
                    END) AS chiffre,
                    SUM(CASE 
                        WHEN xf.C_DocType_ID = 1000646 THEN -1 * TO_NUMBER(mi.QTYENTERED)
                        ELSE TO_NUMBER(mi.QTYENTERED)
                    END) AS qty
                FROM 
                    M_InOut  xf
                JOIN M_INOUTLINE mi ON mi.M_INOUT_ID=xf.M_INOUT_ID
                JOIN C_BPartner cb ON cb.C_BPARTNER_ID=xf.C_BPARTNER_ID
                LEFT JOIN C_InvoiceLine ci ON ci.M_INOUTLINE_ID=mi.M_INOUTLINE_ID
                JOIN M_ATTRIBUTEINSTANCE ma ON ma.M_ATTRIBUTESETINSTANCE_ID=mi.M_ATTRIBUTESETINSTANCE_ID
                JOIN M_PRODUCT m ON m.M_PRODUCT_id=mi.M_PRODUCT_id
                WHERE xf.MOVEMENTDATE BETWEEN TO_DATE(:start_date, 'YYYY-MM-DD') 
                                          AND TO_DATE(:end_date, 'YYYY-MM-DD')
                    AND ma.M_Attribute_ID=1000504
                    AND xf.AD_Org_ID = 1000000
                    AND xf.C_DocType_ID IN (1000013, 1000646)
                    AND xf.M_Warehouse_ID IN (1000724, 1000000, 1000720, 1000725)
                    AND (:fournisseur IS NULL OR UPPER(cb.name) LIKE UPPER(:fournisseur) || '%')
                    AND (:product IS NULL OR UPPER(m.name) LIKE UPPER(:product) || '%')
                GROUP BY ROLLUP(TO_CHAR(xf.MOVEMENTDATE, 'YYYY'), TO_CHAR(xf.MOVEMENTDATE, 'MM'))
                HAVING TO_CHAR(xf.MOVEMENTDATE, 'YYYY') = :year_str OR TO_CHAR(xf.MOVEMENTDATE, 'YYYY') IS NULL
                ORDER BY year NULLS LAST, month NULLS LAST
            """

            params = {
                'start_date': start_date,
                'end_date': end_date,
                'fournisseur': fournisseur or None,
                'product': product or None,
                'year_str': str(year)
            }

            cursor.execute(query, params)
            rows = cursor.fetchall()
            columns = [col[0] for col in cursor.description]
            return [dict(zip(columns, row)) for row in rows]

    except Exception as e:
        logger.error(f"Error fetching fournisseur recap achat data by year: {e}")
        return {"error": "An error occurred while fetching recap achat data by year."}

@app.route('/listfournisseur')
def list_fournisseur():
    try:
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()
            query = """
                SELECT cb.name
                FROM C_BPartner cb
                WHERE cb.AD_Client_ID = 1000000
                  AND cb.ISVENDOR = 'Y'
                  AND cb.ISACTIVE = 'Y'
                ORDER BY cb.name
            """
            cursor.execute(query)
            result = [row[0] for row in cursor.fetchall()]
            return jsonify(result)
    except Exception as e:
        logger.error(f"Error fetching fournisseurs: {e}")
        return jsonify({"error": "Could not fetch fournisseur list"}), 500




@app.route('/listproduct')
def listproduct():
    try:
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()
            query = """
select name from M_PRODUCT
WHERE AD_Client_ID = 1000000
AND AD_Org_ID = 1000000
  AND ISACTIVE = 'Y'
ORDER BY name


            """
            cursor.execute(query)
            result = [row[0] for row in cursor.fetchall()]
            return jsonify(result)
    except Exception as e:
        logger.error(f"Error fetching products: {e}")
        return jsonify({"error": "Could not fetch products list"}), 500
    



@app.route('/fetchZoneClients', methods=['GET'])
def fetch_zone_clients():
    zone = request.args.get('zone')

    if not zone:
        return jsonify({"error": "Missing zone parameter"}), 400

    data = fetch_clients_by_zone(zone)
    return jsonify(data)


def fetch_clients_by_zone(zone):
    try:
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()
            query = """
                SELECT DISTINCT
                    cb.C_BPartner_ID AS "CLIENT_ID",
                    cb.name AS "CLIENT_NAME",
                    sr.name AS "ZONE"
                FROM C_SalesRegion sr
                JOIN C_BPartner_Location bpl ON sr.C_SalesRegion_ID = bpl.C_SalesRegion_ID
                JOIN C_BPartner cb ON cb.C_BPartner_ID = bpl.C_BPartner_ID
                JOIN xx_ca_fournisseur xf ON bpl.C_BPartner_ID = xf.CLIENTID
                WHERE UPPER(sr.name) = UPPER(:zone)
                    AND xf.AD_Org_ID = 1000000
                    AND xf.DOCSTATUS != 'RE'
                ORDER BY cb.name
            """
            
            params = {
                'zone': zone
            }
            
            cursor.execute(query, params)
            rows = cursor.fetchall()
            columns = [col[0] for col in cursor.description]  # Get column names
            return [dict(zip(columns, row)) for row in rows]
    except Exception as e:
        logger.error(f"Error fetching clients by zone: {e}")
        return {"error": "An error occurred while fetching clients by zone."}



@app.route('/listclient')
def list_client():
    try:
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()
            query = """
                SELECT name
                FROM C_BPartner
                WHERE iscustomer = 'Y'
                  AND AD_Client_ID = 1000000
                  AND AD_Org_ID = 1000000
                ORDER BY name
            """
            cursor.execute(query)
            result = [row[0] for row in cursor.fetchall()]
            return jsonify(result)
    except Exception as e:
        logger.error(f"Error fetching clients: {e}")
        return jsonify({"error": "Could not fetch client list"}), 500

@app.route('/listregion')
def list_region():
    try:
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()
            query = """
                SELECT name
                FROM C_SalesRegion
                WHERE ISACTIVE = 'Y'
                  AND AD_Client_ID = 1000000
            """
            cursor.execute(query)
            result = [row[0] for row in cursor.fetchall()]
            return jsonify(result)
    except Exception as e:
        logger.error(f"Error fetching regions: {e}")
        return jsonify({"error": "Could not fetch region list"}), 500








def rot_mont_vente(year_list, fournisseurs, products=None, clients=None, zones=None):
    try:
        # Acquire a connection from the pool
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()

            # Build the supplier condition
            supplier_condition = "1=1"  # Default to include all suppliers
            if fournisseurs:
                if isinstance(fournisseurs, str):
                    fournisseurs = [fournisseurs]

                supplier_conditions = []
                for i, f in enumerate(fournisseurs):
                    bind_var = f'fournisseur_{i}'
                    supplier_conditions.append(f"UPPER(xf.name) LIKE UPPER(:{bind_var}) || '%'")

                supplier_condition = ' OR '.join(supplier_conditions) if supplier_conditions else '1=0'

            # Build product condition for multiple products
            product_condition = "1=1"
            if products:
                if isinstance(products, str):
                    products = [products]
                product_conditions = []
                for i, p in enumerate(products):
                    bind_var = f'product_{i}'
                    product_conditions.append(f"UPPER(xf.product) LIKE UPPER(:{bind_var}) || '%'")
                product_condition = ' OR '.join(product_conditions) if product_conditions else '1=0'

            # Build client condition for multiple clients
            client_condition = "1=1"
            if clients:
                if isinstance(clients, str):
                    clients = [clients]
                client_conditions = []
                for i, c in enumerate(clients):
                    bind_var = f'client_{i}'
                    client_conditions.append(f"UPPER(cb.name) LIKE UPPER(:{bind_var}) || '%'")
                client_condition = ' OR '.join(client_conditions) if client_conditions else '1=0'

            # Build zone condition for multiple zones
            zone_condition = "1=1"
            if zones:
                if isinstance(zones, str):
                    zones = [zones]
                zone_conditions = []
                for i, z in enumerate(zones):
                    bind_var = f'zone_{i}'
                    zone_conditions.append(f"UPPER(sr.name) LIKE UPPER(:{bind_var}) || '%'")
                zone_condition = ' OR '.join(zone_conditions) if zone_conditions else '1=0'

            # Create the years parameter list
            year_placeholders = ', '.join([f':year_{i}' for i in range(len(year_list))])
            
            query = f"""
                WITH monthly_data AS (
                    SELECT 
                        TO_CHAR(xf.MOVEMENTDATE, 'YYYY') AS year,
                        TO_CHAR(xf.MOVEMENTDATE, 'MM') AS month_num,
                        TO_CHAR(xf.MOVEMENTDATE, 'Month') AS month_name,
                        CAST(xf.product AS VARCHAR2(400)) AS produit,
                        xf.name AS fournisseur,
                        cb.name AS client,
                        SUM(xf.TOTALLINE) AS total,
                        SUM(xf.qtyentered) AS qty,
                        CASE 
                            WHEN SUM(xf.CONSOMATION) = 0 THEN 0
                            ELSE ROUND((SUM(xf.TOTALLINE) - SUM(xf.CONSOMATION)) / NULLIF(SUM(xf.CONSOMATION), 0), 4)
                        END AS marge
                    FROM 
                        xx_ca_fournisseur xf
                        LEFT JOIN C_BPartner cb ON cb.C_BPartner_ID = xf.CLIENTID
                        LEFT JOIN AD_User au ON au.AD_User_ID = xf.SALESREP_ID
                        LEFT JOIN C_BPartner_Location bpl ON bpl.C_BPartner_ID = xf.CLIENTID
                        LEFT JOIN C_SalesRegion sr ON sr.C_SalesRegion_ID = bpl.C_SalesRegion_ID
                        LEFT JOIN M_InOut mi ON xf.DOCUMENTNO = mi.DOCUMENTNO
                        LEFT JOIN C_ORDER c ON mi.C_ORDER_ID = c.C_ORDER_ID
                    WHERE 
                        TO_CHAR(xf.MOVEMENTDATE, 'YYYY') IN ({year_placeholders})
                        AND xf.AD_Org_ID = 1000000
                        AND xf.DOCSTATUS != 'RE'
                        AND ({supplier_condition})
                        AND ({product_condition})
                        AND ({client_condition})
                        AND ({zone_condition})
                    GROUP BY 
                        TO_CHAR(xf.MOVEMENTDATE, 'YYYY'),
                        TO_CHAR(xf.MOVEMENTDATE, 'MM'),
                        TO_CHAR(xf.MOVEMENTDATE, 'Month'),
                        xf.product,
                        xf.name,
                        cb.name
                )
                SELECT * FROM monthly_data
                ORDER BY year, month_num, produit
            """

            # Prepare parameters
            params = {}
            
            # Add supplier parameters only if suppliers are specified
            if fournisseurs:
                for i, f in enumerate(fournisseurs):
                    params[f'fournisseur_{i}'] = f

            # Add product parameters
            if products:
                for i, p in enumerate(products):
                    params[f'product_{i}'] = p

            # Add client parameters
            if clients:
                for i, c in enumerate(clients):
                    params[f'client_{i}'] = c

            # Add zone parameters
            if zones:
                for i, z in enumerate(zones):
                    params[f'zone_{i}'] = z

            # Add year parameters
            for i, year in enumerate(year_list):
                params[f'year_{i}'] = str(year)

            # Execute the query with the parameters
            cursor.execute(query, params)

            # Fetch the results
            rows = cursor.fetchall()

            # Format the results into a nested dictionary structure
            result = {}
            
            # Process each row and organize by year and month
            for row in rows:
                year, month_num, month_name, produit, fournisseur, client, total, qty, marge = row
                month_num = month_num.zfill(2)  # Ensure two digits for month number
                
                if year not in result:
                    result[year] = {}
                
                if month_num not in result[year]:
                    result[year][month_num] = {
                        'month_name': month_name.strip(),
                        'details': [],
                        'total': {'QTY': 0, 'TOTAL': 0, 'MARGE': 0}
                    }
                
                # Add product details with supplier and client info
                result[year][month_num]['details'].append({
                    'PRODUIT': produit,
                    'FOURNISSEUR': fournisseur,
                    'CLIENT': client if client else '',
                    'QTY': float(qty) if qty is not None else 0,
                    'TOTAL': float(total) if total is not None else 0,
                    'MARGE': float(marge) if marge is not None else 0
                })
                
                # Update monthly totals
                result[year][month_num]['total']['QTY'] += float(qty) if qty is not None else 0
                result[year][month_num]['total']['TOTAL'] += float(total) if total is not None else 0
                # For marge, we need to recalculate it based on aggregated values
                # This is a simplified approach; you might need to adjust based on your business logic

            # Sort years and months numerically
            sorted_result = {
                str(year): dict(sorted(months.items()))
                for year, months in sorted(result.items())
            }
            return sorted_result

    except Exception as e:
        logger.error(f"Error fetching rotation monthly vente: {e}")
        return {"error": "An error occurred while fetching rotation monthly vente."}

@app.route('/rot_mont_vente', methods=['GET'])
def fetch_rotation_monthly_vente():
    years = request.args.get('years')  # Comma-separated list of years
    fournisseur_param = request.args.get('fournisseur')  # Can be comma-separated
    product_param = request.args.get('product')  # Can be comma-separated
    client_param = request.args.get('client')  # Can be comma-separated
    zone_param = request.args.get('zone')  # Can be comma-separated

    # If years is not provided, use current year
    if not years:
        current_year = datetime.now().year
        years = str(current_year)
    
    # Handle multiple suppliers (comma-separated or single)
    fournisseurs = None
    if fournisseur_param:
        fournisseurs = [f.strip() for f in fournisseur_param.split(',') if f.strip()]

    # Handle multiple products (comma-separated or single)
    products = None
    if product_param:
        products = [p.strip() for p in product_param.split(',') if p.strip()]

    # Handle multiple clients (comma-separated or single)
    clients = None
    if client_param:
        clients = [c.strip() for c in client_param.split(',') if c.strip()]

    # Handle multiple zones (comma-separated or single)
    zones = None
    if zone_param:
        zones = [z.strip() for z in zone_param.split(',') if z.strip()]

    try:
        # Split the years string into a list and convert to integers
        year_list = [int(year.strip()) for year in years.split(',')]
    except ValueError:
        return jsonify({"error": "Invalid years format. Please provide comma-separated years (e.g., 2022,2023,2024)"}), 400

    # Fetch data from the database with the list of parameters
    data = rot_mont_vente(year_list, fournisseurs, products, clients, zones)

    # Return the result as a JSON response
    return jsonify(data)





def fetch_marge_data():
    try:
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()
            query = """
SELECT
    *
FROM
    (
        SELECT
            "source"."FOURNISSEUR" "FOURNISSEUR",
            "source"."PRODUCT" "PRODUCT",
            "source"."P_ACHAT" "P_ACHAT",
            "source"."P_VENTE" "P_VENTE",
            "source"."REM_ACHAT" "REM_ACHAT",
            "source"."REM_VENTE" "REM_VENTE",
            "source"."BON_ACHAT" "BON_ACHAT",
            "source"."BON_VENTE" "BON_VENTE",
            "source"."REMISE_AUTO" "REMISE_AUTO",
            "source"."BONUS_AUTO" "BONUS_AUTO",
            "source"."P_REVIENT" "P_REVIENT",
            "source"."MARGE" "MARGE",
            "source"."LABO" "LABO",
            "source"."PPA" "PPA",
            "source"."LOT" "LOT",
            "source"."QTY" "QTY",
            "source"."GUARANTEEDATE" "GUARANTEEDATE",  -- Added the GUARANTEEDATE column
            "source"."LOCATION" "LOCATION"
        FROM
            (
                SELECT
                    DISTINCT fournisseur,
                    product,
                    p_achat,
                    p_vente,
                    round(rem_achat, 2) AS rem_achat,
                    rem_vente,
                    round(bon_achat, 2) AS bon_achat,
                    bon_vente,
                    remise_auto,
                    bonus_auto,
                    round(p_revient, 2) AS p_revient,
                    LEAST(round((marge), 2), 100) AS marge,
                    labo,
                    lot,
                    qty,
                    guaranteedate,  -- Added the GUARANTEEDATE column
                    CASE 
                        WHEN m_locator_id = 1000614 THEN 'Préparation'
                        WHEN m_locator_id = 1001135 THEN 'HANGAR'
                        WHEN m_locator_id = 1001128 THEN 'Dépot_réserve'
                        WHEN m_locator_id = 1001136 THEN 'HANGAR_'
                        WHEN m_locator_id = 1001020 THEN 'Depot_Vente'
                    END AS location
                FROM
                    (
                        SELECT
                            d.*,
                            LEAST(
                                round(
                                    (((ventef - ((ventef * nvl(rma, 0)) / 100))) - p_revient) / p_revient * 100,
                                    2
                                ), 
                                100
                            ) AS marge
                        FROM
                            (
                                SELECT
                                    det.*,
                                    (det.p_achat - ((det.p_achat * det.rem_achat) / 100)) / (1 + (det.bon_achat / 100)) p_revient,
                                    (
                                        det.p_vente - ((det.p_vente * nvl(det.rem_vente, 0)) / 100)
                                    ) / (
                                        1 + (
                                            CASE
                                                WHEN det.bna > 0 THEN det.bna
                                                ELSE det.bon_vente
                                            END / 100
                                        )
                                    ) ventef
                                FROM
                                    (
                                        SELECT
                                            p.name product,
                                            (
                                                SELECT
                                                    NAME
                                                FROM
                                                    XX_Laboratory
                                                WHERE
                                                    XX_Laboratory_id = p.XX_Laboratory_id
                                            ) labo,
                                            mst.qtyonhand qty,
                                            mst.m_locator_id,
                                            mati.value fournisseur,
                                            mats.guaranteedate,
                                            md.name remise_auto,
                                            sal.description bonus_auto,
                                            md.flatdiscount rma,
                                            TO_NUMBER(
                                                CASE
                                                    WHEN REGEXP_LIKE(sal.name, '^[0-9]+$') THEN sal.name
                                                    ELSE NULL
                                                END
                                            ) AS bna,
                                            (
                                                SELECT
                                                    valuenumber
                                                FROM
                                                    m_attributeinstance
                                                WHERE
                                                    m_attributesetinstance_id = mst.m_attributesetinstance_id
                                                    AND m_attribute_id = 1000501
                                            ) p_achat,
                                            (
                                                SELECT
                                                    valuenumber
                                                FROM
                                                    m_attributeinstance
                                                WHERE
                                                    m_attributesetinstance_id = mst.m_attributesetinstance_id
                                                    AND m_attribute_id = 1001009
                                            ) rem_achat,
                                            (
                                                SELECT
                                                    valuenumber
                                                FROM
                                                    m_attributeinstance
                                                WHERE
                                                    m_attributesetinstance_id = mst.m_attributesetinstance_id
                                                    AND m_attribute_id = 1000808
                                            ) bon_achat,
                                            (
                                                SELECT
                                                    valuenumber
                                                FROM
                                                    m_attributeinstance
                                                WHERE
                                                    m_attributesetinstance_id = mst.m_attributesetinstance_id
                                                    AND m_attribute_id = 1000502
                                            ) p_vente,
                                            (
                                                SELECT
                                                    valuenumber
                                                FROM
                                                    m_attributeinstance
                                                WHERE
                                                    m_attributesetinstance_id = mst.m_attributesetinstance_id
                                                    AND m_attribute_id = 1001408
                                            ) rem_vente,
                                            (
                                                SELECT
                                                    valuenumber
                                                FROM
                                                    m_attributeinstance
                                                WHERE
                                                    m_attributesetinstance_id = mst.m_attributesetinstance_id
                                                    AND m_attribute_id = 1000908
                                            ) bon_vente,
                                            (
                                                SELECT
                                                    valuenumber
                                                FROM
                                                    m_attributeinstance
                                                WHERE
                                                    m_attributesetinstance_id = mst.m_attributesetinstance_id
                                                    AND m_attribute_id = 1000503
                                            ) PPA,
                                            (
                                                SELECT
                                                    lot
                                                FROM
                                                    m_attributesetinstance
                                                WHERE
                                                    m_attributesetinstance_id = mst.m_attributesetinstance_id
                                            ) lot
                                        FROM
                                            m_product p
                                            INNER JOIN m_storage mst ON p.m_product_id = mst.m_product_id
                                            INNER JOIN m_attributeinstance mati ON mst.m_attributesetinstance_id = mati.m_attributesetinstance_id
                                            INNER JOIN m_attributesetinstance mats ON mst.m_attributesetinstance_id = mats.m_attributesetinstance_id
                                            LEFT JOIN C_BPartner_Product cp ON cp.m_product_id = p.m_product_id
                                                OR cp.C_BPartner_Product_id IS NULL
                                            LEFT JOIN M_DiscountSchema md ON cp.M_DiscountSchema_id = md.M_DiscountSchema_id
                                            LEFT JOIN XX_SalesContext sal ON p.XX_SalesContext_ID = sal.XX_SalesContext_ID
                                        WHERE
                                            mati.m_attribute_id = 1000508
                                            AND mst.m_locator_id IN (1001135, 1000614, 1001128, 1001136, 1001020)
                                            AND mst.qtyonhand != 0
                                        ORDER BY
                                            p.name
                                    ) det
                                WHERE
                                    det.rem_achat < 200
                            ) d
                    )
                GROUP BY
                    fournisseur,
                    product,
                    p_achat,
                    p_vente,
                    rem_achat,
                    rem_vente,
                    bon_achat,
                    bon_vente,
                    remise_auto,
                    bonus_auto,
                    p_revient,
                    marge,
                    labo,
                    lot,
                    qty,
                    PPA,
                    guaranteedate,  -- Added to GROUP BY
                    m_locator_id
                ORDER BY
                    fournisseur
            ) "source"
    )
WHERE
    rownum <= 1048575
            """
            cursor.execute(query)
            rows = cursor.fetchall()
            columns = [col[0] for col in cursor.description]  # Get column names
            data = [dict(zip(columns, row)) for row in rows]
            return data
    except Exception as e:
        logger.error(f"Error fetching marge data: {e}")
        return {"error": "An error occurred while fetching marge data."}








if __name__ == "__main__":
    app.run(host='0.0.0.0', port=5001)