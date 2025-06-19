import oracledb
from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
import logging
import pandas as pd
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import datetime

app = Flask(__name__)
CORS(app)  # Enable CORS for all routes

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Configure Oracle database connection pool
DB_POOL = oracledb.create_pool(
    user="compiere",
    password="compiere",
    dsn="192.168.1.213/compiere",
    min=2,
    max=10,
    increment=1
)


def fetch_fournisseur_by_product(product):
    try:
        # Acquire a connection from the pool
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()

            query = """
                SELECT DISTINCT
                    CAST(cb.name AS VARCHAR2(300)) AS FOURNISSEUR
                FROM 
                    M_InOut xf
                JOIN M_INOUTLINE mi ON mi.M_INOUT_ID = xf.M_INOUT_ID
                JOIN C_BPartner cb ON cb.C_BPARTNER_ID = xf.C_BPARTNER_ID
                JOIN M_PRODUCT m ON m.M_PRODUCT_id = mi.M_PRODUCT_id
                WHERE 
                    UPPER(m.name) LIKE UPPER(:product) || '%'
                    AND xf.AD_Org_ID = 1000000
                    AND xf.C_DocType_ID IN (1000013, 1000646)
                    AND xf.M_Warehouse_ID IN (1000724, 1000000, 1000720, 1000725)
                ORDER BY 
                    FOURNISSEUR
            """

            params = {
                'product': product or None
            }

            # Execute the query with the provided parameters
            cursor.execute(query, params)

            # Fetch the results
            rows = cursor.fetchall()

            # Format the results into a simple list of supplier names
            suppliers = [row[0] for row in rows]

            return suppliers
    
    except Exception as e:
        logger.error(f"Error fetching suppliers by product: {e}")
        return {"error": "An error occurred while fetching suppliers by product."}

# Flask route to handle the request
@app.route('/fetchSuppliersByProduct', methods=['GET'])
def fetch_suppliers_by_product():
    product = request.args.get('product')

    # Ensure product is provided
    if not product:
        return jsonify({"error": "Missing product parameter"}), 400

    # Fetch data from the database
    suppliers = fetch_fournisseur_by_product(product)

    # Return the result as a JSON response
    return jsonify(suppliers)




@app.route('/rotation_monthly_achat', methods=['GET'])
def fetch_rotation_achat():
    years = request.args.get('years')  # Comma-separated list of years
    fournisseur_param = request.args.get('fournisseur')  # Can be comma-separated
    product = request.args.get('product')

    # If years is not provided, use current year
    if not years:
        current_year = datetime.now().year
        years = str(current_year)
    
    # Handle multiple suppliers (comma-separated or single)
    if fournisseur_param:
        fournisseurs = [f.strip() for f in fournisseur_param.split(',') if f.strip()]
    else:
        return jsonify({"error": "At least one supplier is required"}), 400

    try:
        # Split the years string into a list and convert to integers
        year_list = [int(year.strip()) for year in years.split(',')]
    except ValueError:
        return jsonify({"error": "Invalid years format. Please provide comma-separated years (e.g., 2022,2023,2024)"}), 400

    # Fetch data from the database with the list of suppliers
    data = fetch_rotation_monthly_achat(year_list, fournisseurs, product)

    # Return the result as a JSON response
    return jsonify(data)


def fetch_rotation_monthly_achat(year_list, fournisseurs, product=None):
    try:
        # Acquire a connection from the pool
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()

            # Build the supplier condition
            if isinstance(fournisseurs, str):
                fournisseurs = [fournisseurs]

            supplier_conditions = []
            for i, f in enumerate(fournisseurs):
                bind_var = f'fournisseur_{i}'
                supplier_conditions.append(f"UPPER(cb.name) LIKE UPPER(:{bind_var}) || '%'")

            supplier_condition = ' OR '.join(supplier_conditions) if supplier_conditions else '1=0'

            # Create the years parameter list
            year_placeholders = ', '.join([f':year_{i}' for i in range(len(year_list))])
            
            query = """
                WITH monthly_data AS (
                    SELECT 
                        TO_CHAR(xf.MOVEMENTDATE, 'YYYY') AS year,
                        TO_CHAR(xf.MOVEMENTDATE, 'MM') AS month_num,
                        TO_CHAR(xf.MOVEMENTDATE, 'Month') AS month_name,
                        m.name AS produit,
                        cb.name AS fournisseur,
                        SUM(CASE 
                            WHEN xf.C_DocType_ID = 1000646 THEN -1 * TO_NUMBER(mi.QTYENTERED)
                            ELSE TO_NUMBER(mi.QTYENTERED)
                        END) AS qty,
                        SUM(CASE 
                            WHEN xf.C_DocType_ID = 1000646 THEN -1 * TO_NUMBER(ma.valuenumber) * TO_NUMBER(mi.QTYENTERED)
                            ELSE TO_NUMBER(ma.valuenumber) * TO_NUMBER(mi.QTYENTERED)
                        END) AS chiffre
                    FROM 
                        M_InOut xf
                        JOIN M_INOUTLINE mi ON mi.M_INOUT_ID = xf.M_INOUT_ID
                        JOIN M_ATTRIBUTEINSTANCE ma ON ma.M_ATTRIBUTESETINSTANCE_ID = mi.M_ATTRIBUTESETINSTANCE_ID
                        JOIN C_BPartner cb ON cb.C_BPARTNER_ID = xf.C_BPARTNER_ID
                        JOIN M_PRODUCT m ON m.M_PRODUCT_id = mi.M_PRODUCT_id
                    WHERE 
                        TO_CHAR(xf.MOVEMENTDATE, 'YYYY') IN (""" + year_placeholders + """)
                        AND xf.AD_Org_ID = 1000000
                        AND xf.C_DocType_ID IN (1000013, 1000646)
                        AND ma.M_Attribute_ID = 1000504
                        AND xf.DOCSTATUS = 'CO'
                        AND xf.M_Warehouse_ID IN (1000724, 1000000, 1000720, 1000725)
                        AND (""" + supplier_condition + """)
                        AND (:product IS NULL OR UPPER(m.name) LIKE UPPER(:product) || '%')
                    GROUP BY 
                        TO_CHAR(xf.MOVEMENTDATE, 'YYYY'),
                        TO_CHAR(xf.MOVEMENTDATE, 'MM'),
                        TO_CHAR(xf.MOVEMENTDATE, 'Month'),
                        m.name,
                        cb.name
                )
                SELECT * FROM monthly_data
                ORDER BY year, month_num, produit
            """

            # Prepare parameters
            params = {'product': product}
            
            # Add supplier parameters
            for i, f in enumerate(fournisseurs):
                params[f'fournisseur_{i}'] = f

            # Add year parameters
            for i, year in enumerate(year_list):
                params[f'year_{i}'] = str(year)

            # Execute the query with the parameters
            cursor.execute(query, params)

            # Fetch the results
            rows = cursor.fetchall()

            # Format the results into a nested dictionary structure
            result = {}
            
            # Process each row and organize by year and month
            for row in rows:
                year, month_num, month_name, produit, fournisseur, qty, chiffre = row
                month_num = month_num.zfill(2)  # Ensure two digits for month number
                
                if year not in result:
                    result[year] = {}
                
                if month_num not in result[year]:
                    result[year][month_num] = {
                        'month_name': month_name.strip(),
                        'details': [],
                        'total': {'QTY': 0, 'CHIFFRE': 0}
                    }
                
                # Add product details with supplier info
                result[year][month_num]['details'].append({
                    'PRODUIT': produit,
                    'FOURNISSEUR': fournisseur,
                    'QTY': float(qty) if qty is not None else 0,
                    'CHIFFRE': float(chiffre) if chiffre is not None else 0
                })
                
                # Update monthly totals
                result[year][month_num]['total']['QTY'] += float(qty) if qty is not None else 0
                result[year][month_num]['total']['CHIFFRE'] += float(chiffre) if chiffre is not None else 0

            # Sort years and months numerically
            sorted_result = {
                str(year): dict(sorted(months.items()))
                for year, months in sorted(result.items())
            }
            return sorted_result

    except Exception as e:
        logger.error(f"Error fetching product recap achat: {e}")
        return {"error": "An error occurred while fetching product recap achat."}


from flask import make_response
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch
from datetime import datetime
import io

# Define a color palette for years (can be extended as needed)
YEAR_COLORS = [
    colors.HexColor('#4E79A7'),  # Blue
    colors.HexColor('#F28E2B'),  # Orange
    colors.HexColor('#E15759'),  # Red
    colors.HexColor('#76B7B2'),  # Teal
    colors.HexColor('#59A14F'),  # Green
    colors.HexColor('#EDC948'),  # Yellow
    colors.HexColor('#B07AA1'),  # Purple
    colors.HexColor('#FF9DA7'),  # Pink
    colors.HexColor('#9C755F'),  # Brown
    colors.HexColor('#BAB0AC')   # Gray
]

@app.route('/rotation_monthly_achat_pdf', methods=['GET'])
def download_product_achat_pdf():
    from flask import request, jsonify

    years = request.args.get('years')
    fournisseur_param = request.args.get('fournisseur')
    product = request.args.get('product')

    if not years:
        years = str(datetime.now().year)

    try:
        year_list = [int(y.strip()) for y in years.split(',')]
    except ValueError:
        return jsonify({"error": "Invalid years format"}), 400

    # Handle multiple suppliers (comma-separated or single)
    if fournisseur_param:
        fournisseurs = [f.strip() for f in fournisseur_param.split(',') if f.strip()]
    else:
        return jsonify({"error": "Fournisseur parameter is required"}), 400

    data = fetch_rotation_monthly_achat(year_list, fournisseurs, product)
    if 'error' in data:
        return jsonify(data), 500

    # Create descriptive filename
    supplier_text = f"{len(fournisseurs)}_suppliers" if len(fournisseurs) > 1 else fournisseurs[0][:20]
    filename = f"achat_recap_{years.replace(',', '-')}_{supplier_text}_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=20, leftMargin=20, topMargin=30, bottomMargin=30)
    elements = []

    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    normal_style = styles["Normal"]
    table_total_style = ParagraphStyle("row", parent=normal_style, alignment=1, textColor=colors.blue, fontName="Helvetica-Bold", fontSize=6)

    table_header_style = ParagraphStyle("header", parent=normal_style, alignment=1, textColor=colors.white, fontSize=8)
    table_row_style = ParagraphStyle("row", parent=normal_style, alignment=1, fontSize=6)

    elements.append(Paragraph("Récapitulatif des Achats", title_style))
    
    # Create supplier display text
    supplier_display = ", ".join(fournisseurs) if len(fournisseurs) <= 3 else f"{len(fournisseurs)} suppliers"
    elements.append(Paragraph(f"Fournisseurs: {supplier_display} | Produit: {product or 'Tous'} | Années: {years}", styles["Heading2"]))
    elements.append(Spacer(1, 0.25 * inch))

    for year_idx, (year, months_data) in enumerate(data.items()):
        year_color = YEAR_COLORS[year_idx % len(YEAR_COLORS)]
        year_style = ParagraphStyle("year", parent=styles["Heading1"], textColor=colors.white, backColor=year_color)
        elements.append(Paragraph(f"Année: {year}", year_style))

        for part_label, month_range in [("Janvier à Juin", range(1, 7)), ("Juillet à Décembre", range(7, 13))]:
            # For multiple suppliers, create a combined product map with supplier info
            if len(fournisseurs) > 1:
                combined_products = {}
                for m in month_range:
                    m_str = str(m).zfill(2)
                    if m_str in months_data:
                        for item in months_data[m_str]["details"]:
                            product_key = f"{item['PRODUIT']} ({item['FOURNISSEUR']})"
                            if product_key not in combined_products:
                                combined_products[product_key] = {}
                            combined_products[product_key][m] = item
                unique_products = list(combined_products.keys())
            else:
                # Single supplier - use original logic
                unique_products = set()
                for m in month_range:
                    m_str = str(m).zfill(2)
                    if m_str in months_data:
                        for item in months_data[m_str]["details"]:
                            unique_products.add(item["PRODUIT"])
                unique_products = list(unique_products)

            if not unique_products:
                continue

            # Header rows
            header_row1 = [Paragraph("Produit", table_header_style)]
            for m in month_range:
                m_str = str(m).zfill(2)
                month_name = months_data.get(m_str, {}).get("month_name", datetime(2025, m, 1).strftime("%b")).capitalize()
                header_row1.extend([Paragraph(month_name, table_header_style), ""])

            header_row2 = [""]
            for _ in month_range:
                header_row2.extend([
                    Paragraph("Qty", table_header_style),
                    Paragraph("Prix", table_header_style)
                ])

            table_data = [header_row1, header_row2]

            # Product rows
            for prod in sorted(unique_products):
                row = [Paragraph(prod, table_row_style)]
                for m in month_range:
                    m_str = str(m).zfill(2)
                    month_data = months_data.get(m_str, {})
                    
                    if len(fournisseurs) > 1:
                        # For combined view, find the specific product entry
                        detail = combined_products.get(prod, {}).get(m)
                        qty = f"{detail['QTY']:,.0f}" if detail else "-"
                        amt = f"{detail['CHIFFRE']:,.2f}" if detail else "-"
                    else:
                        # Single supplier - use original logic
                        detail = next((d for d in month_data.get("details", []) if d["PRODUIT"] == prod), None)
                        qty = f"{detail['QTY']:,.0f}" if detail else "-"
                        amt = f"{detail['CHIFFRE']:,.2f}" if detail else "-"
                    
                    row.append(Paragraph(qty, table_row_style))
                    row.append(Paragraph(amt, table_row_style))
                table_data.append(row)

            # Total row
            # Total row
            total_row = [Paragraph("TOTAL", table_total_style)]
            for m in month_range:
                m_str = str(m).zfill(2)
                month_data = months_data.get(m_str, {})
                total = month_data.get("total", {})

                qty_val = total.get("QTY")
                amt_val = total.get("CHIFFRE")

                qty = f"{qty_val:,.0f}" if qty_val is not None else "-"
                amt = f"{amt_val:,.2f}" if amt_val is not None else "-"

                total_row.append(Paragraph(qty, table_total_style))
                total_row.append(Paragraph(amt, table_total_style))
            table_data.append(total_row)

            # Table construction
            col_widths = [1.9 * inch] + [0.5 * inch] * (len(month_range) * 2)
            table = Table(table_data, colWidths=col_widths)

            # Create span styles for merged headers
            span_style = []
            for i in range(len(month_range)):
                col_start = 1 + i * 2
                span_style.append(('SPAN', (col_start, 0), (col_start + 1, 0)))

            # Final style
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 1), year_color),
                ('TEXTCOLOR', (0, 0), (-1, 1), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 2),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('ROWBACKGROUNDS', (0, 2), (-1, -2), [colors.whitesmoke, colors.white]),
            ] + span_style))

            elements.append(Paragraph(f"{part_label}", styles["Heading3"]))
            elements.append(table)
            elements.append(Spacer(1, 0.3 * inch))

    doc.build(elements)
    buffer.seek(0)
    response = make_response(buffer.getvalue())
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = f"attachment; filename={filename}"

    return response


@app.route('/fetchFournisseurRecapAchatByYear', methods=['GET'])
def fetch_fournisseur_recap_achat_by_year():
    fournisseur = request.args.get('fournisseur')
    product = request.args.get('product')
    years = request.args.getlist('years')  # Get all years parameters
    month = request.args.get('month')

    try:
        years = [int(y) for y in years] if years else None
        month = int(month) if month else None
    except ValueError:
        return jsonify({"error": "Invalid year or month format"}), 400

    if not years:
        # If no years specified, get current year
        current_year = datetime.now().year
        years = [current_year]

    full_data = {}
    for year in years:
        yearly_data = fetch_fournisseur_recap_achat_data_by_year(
            fournisseur=fournisseur,
            product=product,
            year=year,
            month=month
        )
        full_data[str(year)] = yearly_data

    return jsonify(full_data)




def fetch_fournisseur_recap_achat_data_by_year(fournisseur=None, product=None, year=None, month=None):
    try:
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()

            if year and month:
                start_date = f"{year}-{month:02d}-01"
                last_day = 31 if month in [1, 3, 5, 7, 8, 10, 12] else 30
                if month == 2:
                    last_day = 29 if year % 4 == 0 and (year % 100 != 0 or year % 400 == 0) else 28
                end_date = f"{year}-{month:02d}-{last_day}"
            elif year:
                start_date = f"{year}-01-01"
                end_date = f"{year}-12-31"
            else:
                current_year = datetime.now().year
                start_date = f"{current_year}-01-01"
                end_date = f"{current_year}-12-31"
                year = current_year

            query = """
                SELECT 
                    TO_CHAR(xf.MOVEMENTDATE, 'YYYY') AS year,
                    TO_CHAR(xf.MOVEMENTDATE, 'MM') AS month,
                    SUM(CASE 
                        WHEN xf.C_DocType_ID = 1000646 THEN -1 * TO_NUMBER(ma.valuenumber) * TO_NUMBER(mi.QTYENTERED) 
                        ELSE TO_NUMBER(ma.valuenumber) * TO_NUMBER(mi.QTYENTERED) 
                    END) AS chiffre,
                    SUM(CASE 
                        WHEN xf.C_DocType_ID = 1000646 THEN -1 * TO_NUMBER(mi.QTYENTERED)
                        ELSE TO_NUMBER(mi.QTYENTERED)
                    END) AS qty
                FROM 
                    M_InOut  xf
                JOIN M_INOUTLINE mi ON mi.M_INOUT_ID=xf.M_INOUT_ID
                JOIN C_BPartner cb ON cb.C_BPARTNER_ID=xf.C_BPARTNER_ID
                LEFT JOIN C_InvoiceLine ci ON ci.M_INOUTLINE_ID=mi.M_INOUTLINE_ID
                JOIN M_ATTRIBUTEINSTANCE ma ON ma.M_ATTRIBUTESETINSTANCE_ID=mi.M_ATTRIBUTESETINSTANCE_ID
                JOIN M_PRODUCT m ON m.M_PRODUCT_id=mi.M_PRODUCT_id
                WHERE xf.MOVEMENTDATE BETWEEN TO_DATE(:start_date, 'YYYY-MM-DD') 
                                          AND TO_DATE(:end_date, 'YYYY-MM-DD')
                    AND ma.M_Attribute_ID=1000504
                    AND xf.AD_Org_ID = 1000000
                    AND xf.C_DocType_ID IN (1000013, 1000646)
                    AND xf.M_Warehouse_ID IN (1000724, 1000000, 1000720, 1000725)
                    AND (:fournisseur IS NULL OR UPPER(cb.name) LIKE UPPER(:fournisseur) || '%')
                    AND (:product IS NULL OR UPPER(m.name) LIKE UPPER(:product) || '%')
                GROUP BY ROLLUP(TO_CHAR(xf.MOVEMENTDATE, 'YYYY'), TO_CHAR(xf.MOVEMENTDATE, 'MM'))
                HAVING TO_CHAR(xf.MOVEMENTDATE, 'YYYY') = :year_str OR TO_CHAR(xf.MOVEMENTDATE, 'YYYY') IS NULL
                ORDER BY year NULLS LAST, month NULLS LAST
            """

            params = {
                'start_date': start_date,
                'end_date': end_date,
                'fournisseur': fournisseur or None,
                'product': product or None,
                'year_str': str(year)
            }

            cursor.execute(query, params)
            rows = cursor.fetchall()
            columns = [col[0] for col in cursor.description]
            return [dict(zip(columns, row)) for row in rows]

    except Exception as e:
        logger.error(f"Error fetching fournisseur recap achat data by year: {e}")
        return {"error": "An error occurred while fetching recap achat data by year."}

@app.route('/listfournisseur')
def list_fournisseur():
    try:
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()
            query = """
                SELECT cb.name
                FROM C_BPartner cb
                WHERE cb.AD_Client_ID = 1000000
                  AND cb.ISVENDOR = 'Y'
                  AND cb.ISACTIVE = 'Y'
                ORDER BY cb.name
            """
            cursor.execute(query)
            result = [row[0] for row in cursor.fetchall()]
            return jsonify(result)
    except Exception as e:
        logger.error(f"Error fetching fournisseurs: {e}")
        return jsonify({"error": "Could not fetch fournisseur list"}), 500




@app.route('/listproduct')
def listproduct():
    try:
        with DB_POOL.acquire() as connection:
            cursor = connection.cursor()
            query = """
select name from M_PRODUCT
WHERE AD_Client_ID = 1000000
AND AD_Org_ID = 1000000
  AND ISACTIVE = 'Y'
ORDER BY name


            """
            cursor.execute(query)
            result = [row[0] for row in cursor.fetchall()]
            return jsonify(result)
    except Exception as e:
        logger.error(f"Error fetching products: {e}")
        return jsonify({"error": "Could not fetch products list"}), 500
    




if __name__ == "__main__":
    app.run(host='0.0.0.0', port=5001)